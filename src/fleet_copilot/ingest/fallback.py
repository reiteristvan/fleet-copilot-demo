"""A parser that needs no Azure account, and refuses to pretend it is one.

This is what unit tests use, so CI never calls the service. It is deliberately a
weaker parser, not an equivalent one: it produces no heading roles and no table
structure, and it raises on a PDF with no text layer rather than returning an
empty document.

That last refusal is the point. Five of the twenty PDFs in this corpus are
image-only, and one carries a planted prompt injection reachable by no other
route. A fallback that quietly returned nothing for them would keep the suite
green while the OCR path went untested.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
from datetime import UTC, datetime
from typing import Final

from fleet_copilot.ingest.parse import (
    BlockRole,
    ParsedBlock,
    ParsedDocument,
    ParsedPage,
    ParseError,
    ParserId,
)

PDF_CONTENT_TYPE: Final = "application/pdf"
DOCX_CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MODEL_ID: Final = "local"

PIECE_SEPARATOR: Final = "\n\n"
"""Joins extracted pieces. Two characters, counted into every offset after the
first, which is why offsets are accumulated as content is built rather than
searched for afterwards."""


def _pdf_pages(data: bytes) -> list[str]:
    """Extract text per page. Blocking; call via :func:`asyncio.to_thread`."""
    import pymupdf

    # pymupdf.open is an alias for pymupdf.Document, which carries no
    # annotations even though the package ships py.typed.
    with pymupdf.open(  # type: ignore[no-untyped-call]  # pymupdf: unannotated Document
        stream=data, filetype="pdf"
    ) as document:
        return [str(page.get_text()) for page in document]


def _docx_paragraphs(data: bytes) -> list[str]:
    """Extract non-empty paragraph text. Blocking; see above."""
    import docx

    return [
        paragraph.text for paragraph in docx.Document(io.BytesIO(data)).paragraphs if paragraph.text
    ]


def _assemble(doc_id: str, source_sha256: str, pieces: list[tuple[str, int]]) -> ParsedDocument:
    """Build a document from ``(text, page_number)`` pieces, in order.

    Offsets accumulate as the content string is built rather than being searched
    for afterwards: two identical paragraphs on one page -- which handover notes
    produce routinely -- would both find the first occurrence.
    """
    content_parts: list[str] = []
    blocks: list[ParsedBlock] = []
    page_bounds: dict[int, tuple[int, int]] = {}
    cursor = 0

    for text, page_number in pieces:
        if cursor:
            content_parts.append(PIECE_SEPARATOR)
            cursor += len(PIECE_SEPARATOR)
        content_parts.append(text)
        blocks.append(
            ParsedBlock(
                role=BlockRole.PARAGRAPH,
                text=text,
                start=cursor,
                end=cursor + len(text),
                page_number=page_number,
            )
        )
        known = page_bounds.get(page_number)
        page_bounds[page_number] = (cursor if known is None else known[0], cursor + len(text))
        cursor += len(text)

    return ParsedDocument(
        doc_id=doc_id,
        source_sha256=source_sha256,
        parser=ParserId.LOCAL,
        model_id=MODEL_ID,
        api_version=None,
        content="".join(content_parts),
        blocks=tuple(blocks),
        pages=tuple(
            ParsedPage(page_number=number, start=start, end=end)
            for number, (start, end) in sorted(page_bounds.items())
        ),
        parsed_at=datetime.now(UTC),
    )


class LocalParser:
    """Parses PDF and DOCX offline. Implements :class:`DocumentParser`."""

    async def parse(self, data: bytes, *, doc_id: str, content_type: str) -> ParsedDocument:
        """Parse ``data``, or raise :class:`ParseError` saying why it cannot."""
        if content_type == PDF_CONTENT_TYPE:
            pages = await asyncio.to_thread(_pdf_pages, data)
            pieces = [
                (text.strip(), number) for number, text in enumerate(pages, start=1) if text.strip()
            ]
            if not pieces:
                msg = (
                    f"{doc_id}: this PDF has no text layer. Reading it needs OCR, "
                    "which only the Document Intelligence path provides."
                )
                raise ParseError(msg)
        elif content_type == DOCX_CONTENT_TYPE:
            pieces = [(text, 1) for text in await asyncio.to_thread(_docx_paragraphs, data)]
            if not pieces:
                msg = f"{doc_id}: this DOCX contains no paragraph text"
                raise ParseError(msg)
        else:
            msg = (
                f"{doc_id}: cannot parse {content_type!r} locally; "
                f"supported types are {PDF_CONTENT_TYPE} and {DOCX_CONTENT_TYPE}"
            )
            raise ParseError(msg)

        return _assemble(doc_id, hashlib.sha256(data).hexdigest(), pieces)
